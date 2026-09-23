"""구조 → 표준 서식 워드(.docx).

한글(HWP)에서 그대로 열린다. 한글에서 '다른 이름으로 저장 → .hwp' 하면 한글 문서가 된다.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from . import style
from .outline import Block, Report

EMPH_RE = re.compile(r"\*\*(.+?)\*\*")
LABEL_RE = re.compile(r"^(\([^)]{1,12}\))")  # ○ (목적) … 처럼 앞머리 괄호 제목


# ── 저수준 도우미 ──────────────────────────────────────────────

def _font(run, f: style.Font, bold: bool | None = None):
    run.font.size = Pt(f.size)
    run.font.name = f.name
    run.font.bold = f.bold if bold is None else bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), f.name)


def _spacing(p, f: style.Font, before: float | None = None):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(round(f.size * style.LINE_RATIO, 1))
    pf.space_before = Pt(f.before if before is None else before)
    pf.space_after = Pt(f.after)


def _runs(p, text: str, f: style.Font, bold_label: bool = False):
    """**굵게** 와 앞머리 (괄호 제목) 을 굵게 찍는다."""
    if bold_label:
        m = LABEL_RE.match(text)
        if m:
            _font(p.add_run(m.group(1)), f, bold=True)
            text = text[m.end():]
    for i, part in enumerate(EMPH_RE.split(text)):
        if part:
            _font(p.add_run(part), f, bold=bool(i % 2) or f.bold)


def _borders(cell, **sides):
    """sides: top/bottom/left/right = (val, 굵기 1/8pt 단위) 또는 None(없음)."""
    tcpr = cell._element.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        spec = sides.get(side)
        if spec:
            el.set(qn("w:val"), spec[0])
            el.set(qn("w:sz"), str(spec[1]))
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tcpr.append(b)


def _fill(cell, hex_color: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._element.get_or_add_tcPr().append(shd)


def _cell_margins(table, left_right_pt: float, top_bottom_pt: float = 1):
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, pt in (("top", top_bottom_pt), ("left", left_right_pt),
                     ("bottom", top_bottom_pt), ("right", left_right_pt)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(pt * 20)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f = style.Font(style.MYEONGJO, 11)
    _font(p.add_run("- "), f)
    run = p.add_run()
    _font(run, f)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)
    _font(p.add_run(" -"), f)


def _box(doc, text: str, f: style.Font, align, borders: dict, pad: float):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    _borders(cell, **borders)
    p = cell.paragraphs[0]
    p.alignment = align
    _spacing(p, f, before=pad)
    p.paragraph_format.space_after = Pt(pad)
    _runs(p, text, f)
    return t


# ── 블록별 ────────────────────────────────────────────────────

def _indent(p, level: int, f: style.Font):
    """n칸 띄움 + 내어쓰기: 두 번째 줄이 기호 다음 글자에 맞춰 떨어지게."""
    space = f.size / 2          # 공백 한 칸 ≈ 글자 반 폭
    mark = f.size + space       # 기호 한 글자 + 뒤 공백
    pf = p.paragraph_format
    pf.left_indent = Pt(level * space + mark)
    pf.first_line_indent = Pt(-mark)


def _bullet(doc, b: Block):
    mark = style.BULLET_ALIASES.get(b.mark, b.mark)
    if mark in style.BULLETS:
        _, spaces, f = style.BULLETS[mark]
    else:  # ☞ ⇒ → 는 '-' 깊이
        _, spaces, f = style.BULLETS["-"]
    p = doc.add_paragraph()
    _spacing(p, f)
    _indent(p, spaces, f)
    _font(p.add_run(f"{mark} "), f)
    _runs(p, b.text, f, bold_label=(mark == "○"))


def _heading(doc, b: Block, first: bool):
    f = style.HEADING
    p = doc.add_paragraph()
    _spacing(p, f, before=f.before / 2 if first else None)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(f"{b.mark} "), f)
    _runs(p, b.text, f)


def _roman(doc, b: Block, first: bool):
    f = style.ROMAN
    p = doc.add_paragraph()
    _spacing(p, f, before=f.before / 2 if first else None)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(f"{b.mark}. "), f)
    _runs(p, b.text, f)


def _text_width(s: str) -> int:
    """한글·한자 2칸, 나머지 1칸으로 친 글자 폭."""
    s = EMPH_RE.sub(r"\1", s)
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)


def _col_widths(rows, total_mm: float) -> list[float]:
    """가장 긴 칸에 비례해 나누되, 좁은 칸도 최소 폭은 준다."""
    need = [max(_text_width(r[c]) for r in rows) + 2 for c in range(len(rows[0]))]
    floor = total_mm / len(need) * 0.45
    raw = [total_mm * n / sum(need) for n in need]
    short = sum(max(0, floor - w) for w in raw)
    spare = sum(w - floor for w in raw if w > floor) or 1
    return [floor if w <= floor else w - short * (w - floor) / spare for w in raw]


def _row_flags(row, header: bool):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))       # 한 행이 두 쪽에 걸치지 않게
    if header:
        trpr.append(OxmlElement("w:tblHeader"))   # 쪽이 넘어가면 머리행 반복


def _table(doc, b: Block, total_mm: float):
    f = style.TABLE
    rows = b.rows
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _cell_margins(t, style.TABLE_CELL_MARGIN_PT)
    widths = _col_widths(rows, total_mm)
    # 칸 너비는 표 격자(gridCol)에도 적어야 한글·리브레오피스가 따른다
    for col, w in zip(t._tbl.tblGrid.iterchildren(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w / 25.4 * 1440)))
    for r, row in enumerate(rows):
        _row_flags(t.rows[r], header=(r == 0))
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Mm(widths[c])
            p = cell.paragraphs[0]
            # 짧은 표(12행 이하)는 통째로 한 쪽에 두려고 다음 행과 붙인다
            p.paragraph_format.keep_with_next = len(rows) <= 12 and r < len(rows) - 1
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _spacing(p, f, before=0)
            _runs(p, val, f)
            if r == 0:
                _fill(cell, style.TABLE_HEADER_FILL)
                for run in p.runs:
                    run.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(rep: Report, path: str | Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    P = style.PAGE
    sec.page_width, sec.page_height = Mm(P["width_mm"]), Mm(P["height_mm"])
    sec.top_margin, sec.bottom_margin = Mm(P["top_mm"]), Mm(P["bottom_mm"])
    sec.left_margin, sec.right_margin = Mm(P["left_mm"]), Mm(P["right_mm"])
    sec.header_distance, sec.footer_distance = Mm(P["header_mm"]), Mm(P["footer_mm"])
    _page_number(sec)

    # 새 문서에 기본으로 있는 빈 문단을 제목 상자 앞에서 지운다
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)

    # 제목 상자
    _box(doc, rep.title, style.TITLE, WD_ALIGN_PARAGRAPH.CENTER,
         {s: ("single", 12) for s in ("top", "bottom", "left", "right")}, pad=8)

    if rep.date or rep.dept:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _spacing(p, style.META)
        _runs(p, "< " + ", ".join(x for x in (rep.date, rep.dept) if x) + " >", style.META)

    # 개요 상자 (위·아래 겹줄)
    if rep.summary:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        _box(doc, rep.summary, style.SUMMARY, WD_ALIGN_PARAGRAPH.LEFT,
             {"top": ("double", 6), "bottom": ("double", 6)}, pad=4)

    first = True
    for b in rep.blocks:
        if b.kind == "heading":
            _heading(doc, b, first)
        elif b.kind == "roman":
            _roman(doc, b, first)
        elif b.kind == "bullet":
            _bullet(doc, b)
        elif b.kind == "table":
            _table(doc, b, P["width_mm"] - P["left_mm"] - P["right_mm"])
        elif b.kind == "attach":
            p = doc.add_paragraph()
            _spacing(p, style.BODY, before=20)
            tail = "" if b.text.rstrip().endswith("끝.") else "  끝."
            _runs(p, f"붙임  {b.text}{tail}", style.BODY)
        else:
            p = doc.add_paragraph()
            _spacing(p, style.BODY)
            _runs(p, b.text, style.BODY)
        first = False

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
